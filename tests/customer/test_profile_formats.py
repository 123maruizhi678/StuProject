"""多工作表、Word、旧版 Excel 及否定/冲突语义回归；只用合成资料。"""
import io
from zipfile import ZipFile

import openpyxl
import pytest
import xlwt

from app.clients.customer.extract_dify import ExtractDifyClient
from app.services.customer.errors import BadRequestError
from app.services.customer.profile_input import extract_profile, preview_input
from app.services.customer.profile_rules import evaluate_profile


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.setattr(ExtractDifyClient, 'extract_fields', lambda *a, **kw: None)


def test_word_paragraphs_and_table_cells_are_extracted():
    content = '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
    <w:p><w:r><w:t>姓名：合成客户；年龄：18；学历：高中毕业</w:t></w:r></w:p>
    <w:tbl><w:tr><w:tc><w:p><w:r><w:t>意向国家：新加坡；留学意愿：是</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
    <w:p><w:r><w:t>升学需求：是；家庭支持：是</w:t></w:r></w:p></w:body></w:document>'''
    out = io.BytesIO()
    with ZipFile(out, 'w') as archive:
        archive.writestr('word/document.xml', content)
    result = preview_input(file_name='简历.docx', file_bytes=out.getvalue())
    assert result['items'][0]['facts']['full_name'] == '合成客户'
    assert result['items'][0]['evaluation']['products'][0]['status'] == 'matched'


@pytest.mark.parametrize('extension', ['xlsx', 'xls'])
def test_all_worksheets_are_read_without_mixing_customers(extension):
    rows = [['姓名', '年龄', '学历'], ['合成甲', 18, '高中毕业']]
    out = io.BytesIO()
    if extension == 'xlsx':
        book = openpyxl.Workbook()
        book.remove(book.active)
        for title, age in [('一班', 18), ('二班', 36)]:
            sheet = book.create_sheet(title)
            sheet.append([None, None, None])
            sheet.append(rows[0])
            sheet.append([title, age, '高中毕业'])
        notes = book.create_sheet('填写说明')
        notes.append(['请一人一行'])
        book.save(out)
        book.close()
    else:
        book = xlwt.Workbook()
        for title, age in [('一班', 18), ('二班', 36)]:
            sheet = book.add_sheet(title)
            for r, row in enumerate([rows[0], [title, age, '高中毕业']], 1):
                for c, value in enumerate(row):
                    sheet.write(r, c, value)
        book.save(out)
    result = preview_input(file_name='合成客户.' + extension, file_bytes=out.getvalue())
    assert result['total'] == 2
    assert [(i['sheet_name'], i['row_number'], i['facts']['age']) for i in result['items']] == [('一班', 3, 18), ('二班', 3, 36)]
    assert result['items'][0]['evaluation']['products'][1]['status'] == 'pending'
    assert result['items'][1]['evaluation']['products'][1]['status'] == 'not_matched'
    if extension == 'xlsx':
        assert '填写说明' in result['items'][0]['warnings'][0]


def test_country_list_with_comma_and_equivalent_boolean_labels():
    facts, warnings = extract_profile('姓名：合成甲，年龄：24，意向国家：新加坡,德国，家庭支持：是；家庭是否支持：true')
    assert facts['preferred_countries'] == ['新加坡', '德国']
    assert facts['family_support'] is True and not warnings


@pytest.mark.parametrize('text,field', [
    ('愿意学习德语，但另一处记录不愿意学习德语', 'german_learning_willingness'),
    ('家庭支持：是；家人不支持', 'family_support'),
    ('家人支持，但家长不支持', 'family_support'),
])
def test_contradictions_stay_unknown(text, field):
    facts, warnings = extract_profile('姓名：合成甲；' + text)
    assert field not in facts and warnings


@pytest.mark.parametrize('text', ['计划考取德语B1', '没有德语B1', '德语B1未通过', '目标德语B1'])
def test_language_targets_or_failed_tests_are_not_current_qualifications(text):
    facts, _ = extract_profile('姓名：合成甲；' + text)
    assert 'german_level' not in facts


def test_explicit_natural_language_intent_is_understood():
    facts, _ = extract_profile('姓名：合成甲；24岁；本科毕业；计划赴德国工作；希望就业；不希望提升学历；不想找工作')
    assert facts['study_abroad_willingness'] is True
    assert facts['preferred_countries'] == ['德国']
    assert facts['education_goal'] is False
    assert 'employment_goal' not in facts  # 两处就业意愿矛盾，需要核实。


def test_supportive_refusal_is_visible_in_review_questions():
    product = evaluate_profile({'family_support': False})['products'][0]
    assert product['status'] == 'pending'
    assert any(i['field'] == 'family_support' for i in product['missing_information'])


def test_utf16_tsv_is_supported():
    result = preview_input(file_name='客户.tsv', file_bytes='姓名\t年龄\t学历\n合成甲\t24\t本科'.encode('utf-16'))
    assert result['items'][0]['facts']['age'] == 24


def test_malformed_word_does_not_return_empty_success():
    with pytest.raises(BadRequestError, match='Word'):
        preview_input(file_name='客户.docx', file_bytes=b'invalid')


def test_word_table_with_separate_key_and_value_cells():
    out = io.BytesIO()
    with ZipFile(out, 'w') as archive:
        archive.writestr('word/document.xml', '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:tbl><w:tr>
        <w:tc><w:p><w:r><w:t>姓名</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>合成甲</w:t></w:r></w:p></w:tc>
        <w:tc><w:p><w:r><w:t>年龄</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>24</w:t></w:r></w:p></w:tc>
        </w:tr></w:tbl></w:body></w:document>''')
    facts = preview_input(file_name='客户.docx', file_bytes=out.getvalue())['items'][0]['facts']
    assert facts == {'full_name': '合成甲', 'age': 24}


def test_multiple_named_customers_are_not_combined_into_one_profile():
    with pytest.raises(BadRequestError, match='多个不同客户'):
        preview_input(text='姓名：合成甲；年龄：18；姓名：合成乙；年龄：24')


def test_contradictory_study_intent_is_not_restored_by_country_inference():
    facts, warnings = extract_profile('姓名：合成甲；不愿意出国，但另一处写愿意赴德国工作')
    assert 'study_abroad_willingness' not in facts and warnings


def test_negated_destination_is_not_positive_study_intent():
    facts, _ = extract_profile('姓名：合成甲；不希望去德国工作')
    assert facts.get('study_abroad_willingness') is not True


def test_workbook_customer_limit_applies_across_sheets():
    book = openpyxl.Workbook()
    for sheet in [book.active, book.create_sheet('第二表')]:
        sheet.append(['姓名'])
        for i in range(101):
            sheet.append(['合成' + str(i)])
    out = io.BytesIO()
    book.save(out)
    book.close()
    with pytest.raises(BadRequestError, match='200'):
        preview_input(file_name='客户.xlsx', file_bytes=out.getvalue())
